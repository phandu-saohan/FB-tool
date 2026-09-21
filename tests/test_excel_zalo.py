import unittest
import io
import openpyxl
from app.services.email.excel_service import ExcelContactService
from app.database.models import EmailCampaignRecipient


class TestExcelZaloService(unittest.TestCase):
    def test_normalize_phone(self):
        self.assertEqual(ExcelContactService.normalize_phone("0912 345 678"), "0912345678")
        self.assertEqual(ExcelContactService.normalize_phone("+84912345678"), "0912345678")
        self.assertEqual(ExcelContactService.normalize_phone("84912345678"), "0912345678")
        self.assertEqual(ExcelContactService.normalize_phone("0912.345.678"), "0912345678")
        self.assertEqual(ExcelContactService.normalize_phone("0912-345-678"), "0912345678")
        self.assertIsNone(ExcelContactService.normalize_phone("invalid"))

    def test_to_zalo_oa_format(self):
        self.assertEqual(ExcelContactService.to_zalo_oa_format("0912345678"), "84912345678")
        self.assertEqual(ExcelContactService.to_zalo_oa_format("+84912345678"), "84912345678")
        self.assertEqual(ExcelContactService.to_zalo_oa_format("0987654321"), "84987654321")

    def test_sample_template_generation(self):
        buf = ExcelContactService.generate_sample_template()
        self.assertGreater(buf.getbuffer().nbytes, 100)
        wb = openpyxl.load_workbook(buf)
        ws = wb.active
        self.assertEqual(ws.cell(row=1, column=1).value, "Tên")
        self.assertEqual(ws.cell(row=1, column=2).value, "Email")
        self.assertEqual(ws.cell(row=1, column=3).value, "Số điện thoại")

    def test_parse_excel_file(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Tên", "Email", "Số điện thoại"])
        ws.append(["Nguyễn Văn A", "a@test.com", "0912345678"])
        ws.append(["Trần Thị B", "b@test.com", "+84988888888"])
        ws.append(["Lê C", "invalid-email", "0900000000"])  # invalid email
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        valid, errors = ExcelContactService.parse_contacts_file(buf.getvalue(), "test.xlsx")
        self.assertEqual(len(valid), 2)
        self.assertEqual(len(errors), 1)
        self.assertEqual(valid[0]["email"], "a@test.com")
        self.assertEqual(valid[0]["phone"], "0912345678")
        self.assertEqual(valid[1]["phone"], "0988888888")

    def test_export_zalo_oa(self):
        r1 = EmailCampaignRecipient(id=1, email="a@test.com", name="A", phone="0912345678")
        r2 = EmailCampaignRecipient(id=2, email="b@test.com", name="B", phone="0988888888")
        buf = ExcelContactService.export_zalo_oa_contacts("Campaign 1", [r1, r2])
        self.assertGreater(buf.getbuffer().nbytes, 100)
        wb = openpyxl.load_workbook(buf)
        ws = wb.active
        self.assertEqual(ws.cell(row=2, column=2).value, "84912345678")
        self.assertEqual(ws.cell(row=3, column=2).value, "84988888888")

    def test_banner_title_and_header_on_row_2(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["DANH SÁCH KHÁCH HÀNG VIP THÁNG 9/2026"])  # Banner row 1
        ws.append(["STT", "Họ và tên", "Số điện thoại", "Email", "Ghi chú"])  # Header row 2
        ws.append([1, "Bác sĩ Hùng", "0912334455", "dr.hung@clinic.vn", "Ưu tiên"])
        ws.append([2, "Dược sĩ Linh", "0987654321", "linh.duoc@phar.vn", "Đã liên hệ"])
        ws.append(["Tổng cộng", "2 khách hàng", "", "", ""])  # Footer summary row

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        contacts, errors = ExcelContactService.parse_contacts_file(buf.getvalue(), "banner.xlsx")
        self.assertEqual(len(contacts), 2)
        self.assertEqual(contacts[0]["email"], "dr.hung@clinic.vn")
        self.assertEqual(contacts[0]["name"], "Bác sĩ Hùng")
        self.assertEqual(contacts[0]["phone"], "0912334455")
        self.assertEqual(contacts[0]["zalo_phone"], "84912334455")
        self.assertEqual(contacts[1]["email"], "linh.duoc@phar.vn")

    def test_reversed_columns_and_no_headers(self):
        # Raw file: Email | Phone | Name (No header row)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["alpha@domain.com", "0901234567", "Nguyễn Văn Alpha"])
        ws.append(["beta@domain.com", "0907654321", "Trần Thị Beta"])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        contacts, errors = ExcelContactService.parse_contacts_file(buf.getvalue(), "noheader.xlsx")
        self.assertEqual(len(contacts), 2)
        self.assertEqual(contacts[0]["email"], "alpha@domain.com")
        self.assertEqual(contacts[0]["phone"], "0901234567")
        self.assertEqual(contacts[0]["name"], "Nguyễn Văn Alpha")

    def test_semicolon_csv(self):
        csv_content = (
            "Họ tên;Email;SĐT\n"
            "Vũ Minh;vuminh@yahoo.com;0933112233\n"
            "Hoàng Nam;hoangnam@gmail.com;+84944556677\n"
        ).encode('utf-8-sig')

        contacts, errors = ExcelContactService.parse_contacts_file(csv_content, "danh_sach.csv")
        self.assertEqual(len(contacts), 2)
        self.assertEqual(contacts[0]["email"], "vuminh@yahoo.com")
        self.assertEqual(contacts[0]["name"], "Vũ Minh")
        self.assertEqual(contacts[0]["phone"], "0933112233")
        self.assertEqual(contacts[1]["phone"], "0944556677")
        self.assertEqual(contacts[1]["zalo_phone"], "84944556677")


if __name__ == "__main__":
    unittest.main()

